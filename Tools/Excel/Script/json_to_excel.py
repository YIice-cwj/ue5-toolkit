#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
JSON 反写 Excel 数据表同步工具 (json_to_excel.py) —— 双向数据同步（企业级架构强化版）

核心特性：
1. 单一真源（table_schema.py 契约）：三行表头契约与类型编解码完全对齐；
2. 跨工作表安全：基于 table_schema.resolve_openpyxl_sheet 精确定位目标 Sheet，杜绝选表分裂；
3. 写前快照与原子落盘：改动前复制原文件到 DataTables_Backup 并落清单，写入 .tmp 后 os.replace 原子替换；
4. 内存前置数据校验：落盘前检查主键非空与类型可转性，防坏数据落盘；
5. 空值语义保护：Excel 留空单元格默认保留不覆盖，支持 --fill-defaults 显式回填零值；
6. 公式盾牌（Formula Shield）：消费配置 back_write.protect_formulas，'=' 开头单元格默认保护；
7. 行生命周期安全：精准定位数据区末行插行（insert_rows）克隆样式；删除行时对公式表触发强告警；
8. 语义对拍：浮点容差参数化（--float-epsilon），数组按元素序列等价互认；
9. 完整回滚支持：--rollback [--snapshot <目录>] 依据清单无损还原。

用法：
  python3 json_to_excel.py [--config <path>] [--table <表名>] [--dry-run] [--allow-delete]
                           [--force-formula-overwrite] [--fill-defaults] [--float-epsilon 1e-6]
                           [--no-backup] [--strict]
  python3 json_to_excel.py [--config <path>] --rollback [--snapshot <快照目录>] [--table <表名>]
"""

import argparse
from copy import copy
import dataclasses
from datetime import datetime
import hashlib
import json
import os
import re
import shutil
import sys
import zipfile
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

import table_rules
import table_schema
import tool_config
from naming_utils import to_property_name

# 备份快照目录名格式（YYYYMMDD_HHMMSS）
SNAPSHOT_DIR_PATTERN = re.compile(r"^\d{8}_\d{6}$")

# 快照清单文件名与临时文件后缀
MANIFEST_FILENAME = "_manifest.json"
TMP_SUFFIX = ".tmp"


@dataclasses.dataclass
class SyncResult:
    """单表同步结果（含差异、告警、转换失败、规则问题与错误）"""
    has_changes: bool = False
    diff_list: list = dataclasses.field(default_factory=list)
    warnings: list = dataclasses.field(default_factory=list)
    conversion_errors: list = dataclasses.field(default_factory=list)
    drift_issues: list = dataclasses.field(default_factory=list)   # JSON 与 Excel 表头结构漂移（多余字段/缺列）
    rule_issues: list = dataclasses.field(default_factory=list)   # 写前规则校验问题（table_rules）
    text_literal_count: int = 0                                   # JSON 中 FText 字面量处数（非派生 JSON 信号）
    backup_path: str = None
    error: str = None


def file_sha1(path: str) -> str:
    """计算文件 sha1（备份清单与验证用）"""
    digest = hashlib.sha1()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_cell_style(src_cell: Any, dst_cell: Any) -> None:
    """从源单元格完整深拷贝样式至目标单元格（字体、填充、边框、对齐、数字格式）"""
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)
        dst_cell.alignment = copy(src_cell.alignment)


def rotate_backups(backup_root: str, max_backups: int) -> None:
    """清理多余的历史快照，保证最多保留 max_backups 份最新快照"""
    if not os.path.isdir(backup_root) or max_backups <= 0:
        return

    snapshots = [
        name for name in os.listdir(backup_root)
        if os.path.isdir(os.path.join(backup_root, name)) and SNAPSHOT_DIR_PATTERN.match(name)
    ]
    snapshots.sort()

    while len(snapshots) > max_backups:
        oldest = snapshots.pop(0)
        oldest_path = os.path.join(backup_root, oldest)
        try:
            shutil.rmtree(oldest_path)
            print(f"[备份轮转] 已自动清理历史快照: {oldest}")
        except OSError as e:
            print(f"[警告] 无法清理过期备份目录 {oldest_path}: {e}", file=sys.stderr)


class BackupSession:
    """写前快照会话：惰性创建时间戳目录，逐表备份“原文件”并落 sha1 清单"""

    def __init__(self, backup_root: str, excel_dir: str, max_backups: int = 10) -> None:
        self.backup_root = backup_root
        self.excel_dir = excel_dir
        self.max_backups = max_backups
        self.snapshot_dir = None
        self.entries: List[Dict[str, str]] = []

    def backup(self, xlsx_path: str) -> str:
        """
        在写入前备份指定 Excel 原文件
        @param xlsx_path 待写入的 Excel 绝对路径（调用时尚未被修改）
        @return 备份文件路径
        """
        if self.snapshot_dir is None:
            self.snapshot_dir = os.path.join(self.backup_root, datetime.now().strftime("%Y%m%d_%H%M%S"))

        rel_path = os.path.relpath(xlsx_path, self.excel_dir)
        dst_path = os.path.join(self.snapshot_dir, rel_path)
        os.makedirs(os.path.dirname(dst_path), exist_ok=True)
        shutil.copy2(xlsx_path, dst_path)
        self.entries.append({"rel_path": rel_path, "sha1": file_sha1(dst_path)})
        return dst_path

    def finalize(self) -> Optional[str]:
        """落快照清单并轮转历史快照；无备份条目返回 None"""
        if self.snapshot_dir is None or not self.entries:
            return None

        manifest = {
            "timestamp": os.path.basename(self.snapshot_dir),
            "excel_dir": self.excel_dir,
            "entries": self.entries,
        }
        manifest_path = os.path.join(self.snapshot_dir, MANIFEST_FILENAME)
        with open(manifest_path, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(manifest, handle, ensure_ascii=False, indent=2)

        rotate_backups(self.backup_root, self.max_backups)
        return self.snapshot_dir


def resolve_backup_dir(cfg: Dict[str, Any], cfg_path: Optional[str]) -> str:
    """解析备份根目录：配置 back_write.backup_dir 优先，否则回落工具目录旁的 DataTables_Backup"""
    configured = cfg.get("backup_dir")
    if configured:
        return configured
    base = os.path.dirname(os.path.abspath(cfg_path)) if cfg_path else os.getcwd()
    return os.path.normpath(os.path.join(base, "../../DataTables_Backup"))


def resolve_snapshot_dir(backup_root: str, snapshot: Optional[str] = None) -> Optional[str]:
    """解析待回滚快照目录：显式指定优先，否则取最新时间戳快照"""
    if snapshot:
        return snapshot if os.path.isdir(snapshot) else None
    if not os.path.isdir(backup_root):
        return None

    snapshots = sorted(
        name for name in os.listdir(backup_root)
        if os.path.isdir(os.path.join(backup_root, name)) and SNAPSHOT_DIR_PATTERN.match(name)
    )
    if not snapshots:
        return None
    return os.path.join(backup_root, snapshots[-1])


def perform_rollback(backup_root: str, snapshot: Optional[str] = None, table: Optional[str] = None) -> Tuple[bool, str]:
    """
    按快照清单把写前文件原样还原回 Excel（只复制，不删除任何文件）
    @return (是否成功, 结果描述)
    """
    snapshot_dir = resolve_snapshot_dir(backup_root, snapshot)
    if snapshot_dir is None:
        return False, f"未找到可用快照目录（备份根: {backup_root}）"

    manifest_path = os.path.join(snapshot_dir, MANIFEST_FILENAME)
    if not os.path.isfile(manifest_path):
        return False, f"快照清单缺失（该快照可能由旧版本产生）: {manifest_path}"

    with open(manifest_path, "r", encoding="utf-8") as handle:
        manifest = json.load(handle)

    excel_dir = manifest.get("excel_dir") or ""
    if not excel_dir:
        return False, f"快照清单缺少 excel_dir 字段: {manifest_path}"

    restored = []
    skipped = []
    for entry in manifest.get("entries", []):
        rel_path = entry.get("rel_path") or ""
        if not rel_path:
            continue
        if table:
            table_name = os.path.splitext(os.path.basename(rel_path))[0]
            match_target = table.replace("DT_", "").strip()
            if table != table_name and match_target != table_name.replace("DT_", "").strip():
                skipped.append(rel_path)
                continue

        src_path = os.path.join(snapshot_dir, rel_path)
        dst_path = os.path.join(excel_dir, rel_path)
        if not os.path.isfile(src_path):
            skipped.append(rel_path)
            continue

        shutil.copy2(src_path, dst_path)
        restored.append(rel_path)

    return True, f"已从快照 {os.path.basename(snapshot_dir)} 还原 {len(restored)} 张表（跳过 {len(skipped)}）"


def perform_sync_for_table(xlsx_path: str, json_path: str, options: Any, backup_session: Optional[BackupSession] = None) -> SyncResult:
    """
    针对单张数据表执行 JSON ➔ Excel 反向同步（写前备份 + 临时文件原子替换）
    @param xlsx_path 目标 Excel 路径
    @param json_path 源 JSON 路径
    @param options 命令行与配置选项
    @param backup_session 写前快照会话（None 表示不备份）
    @return SyncResult

    实现思路：读取 JSON（解析失败立即报错）→ 打开 Excel 并统一选表 → 解析表头三行契约 →
    扫描已有数据行并做写前内存校验（类型可转性 / 结构漂移 / 非派生探测 / 主键规则）→
    比对更新与新增行（公式保护、空值保护、语义对拍）→ 物理删行（公式表告警）→
    最后走「写前备份 → 写临时文件 → 原子替换」的落盘安全链路。
    """
    result = SyncResult()
    table_name = os.path.splitext(os.path.basename(xlsx_path))[0]
    float_eps = getattr(options, "float_epsilon", table_schema.DEFAULT_FLOAT_EPSILON)
    fill_defaults = getattr(options, "fill_defaults", False)

    protect_formulas = not getattr(options, "force_formula", False)
    if hasattr(options, "protect_formulas_config") and not getattr(options, "force_formula", False):
        protect_formulas = bool(options.protect_formulas_config)

    try:
        with open(json_path, "r", encoding="utf-8") as f:
            json_records = json.load(f)
    except (OSError, json.JSONDecodeError, UnicodeDecodeError) as e:
        result.error = f"读取 JSON 失败 {json_path}: {e}"
        return result

    if not isinstance(json_records, list):
        result.error = f"JSON 根结构必须为对象数组: {json_path}"
        return result

    json_map = {}
    json_keys_ordered = []
    for record in json_records:
        if not isinstance(record, dict):
            continue
        key = record.get("RowName") or record.get("Name")
        if key is not None and str(key).strip():
            key_str = str(key).strip()
            json_map[key_str] = record
            json_keys_ordered.append(key_str)

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    except (OSError, zipfile.BadZipFile, InvalidFileException) as e:
        result.error = f"打开 Excel 失败 {xlsx_path}: {e}"
        return result

    sheet, sheet_title = table_schema.resolve_openpyxl_sheet(wb, table_name)
    if len(wb.sheetnames) > 1 and sheet_title != wb.active.title:
        result.warnings.append(f"多工作表匹配：已锁定数据表工作表 '{sheet_title}'（当前活动工作表为 '{wb.active.title}'）")

    if sheet.max_row < 3:
        result.error = f"Excel 表头不足 3 行契约: {xlsx_path} (工作表: {sheet_title})"
        return result

    row1 = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    row2 = [sheet.cell(row=2, column=c).value for c in range(1, sheet.max_column + 1)]
    row3 = [sheet.cell(row=3, column=c).value for c in range(1, sheet.max_column + 1)]

    columns = table_schema.parse_header_columns(row1, row2, row3)
    if not columns:
        result.error = f"未检测到有效字段列: {xlsx_path}"
        return result

    row_name_col = next((c for c in columns if c.is_key), columns[0])
    valid_cols_by_name = {c.field_name: c for c in columns}

    excel_rows = {}
    has_formula_in_data = False
    for r in range(4, sheet.max_row + 1):
        cell_val = sheet.cell(row=r, column=row_name_col.col_idx + 1).value
        if cell_val is not None and str(cell_val).strip():
            excel_rows[str(cell_val).strip()] = r
        for col in columns:
            c_val = sheet.cell(row=r, column=col.col_idx + 1).value
            if c_val is not None and str(c_val).strip().startswith("="):
                has_formula_in_data = True

    invalid_cells = set()
    for key, record in json_map.items():
        for col in columns:
            prop_name = to_property_name(col.field_name, col.cpp_type)
            if col.field_name in record:
                raw_val = record[col.field_name]
            elif prop_name in record:
                raw_val = record[prop_name]
            elif col.field_name == "RowName" and "Name" in record:
                raw_val = record["Name"]
            else:
                continue

            try:
                table_schema.convert_for_excel(raw_val, col.cpp_type, fill_defaults=fill_defaults, strict=True)
            except table_schema.SchemaValidationError as e:
                result.conversion_errors.append(f"[{key}.{col.field_name}] {e}")
                invalid_cells.add((key, col.field_name))

    record_keys = set()
    for record in json_map.values():
        record_keys |= set(record.keys())
    column_names = {c.field_name for c in columns}
    recognized_names = set(column_names)
    for c in columns:
        recognized_names.add(to_property_name(c.field_name, c.cpp_type))
    recognized_names |= {"Name", "RowName"}
    json_extra_fields = sorted(record_keys - recognized_names)
    matched_excel_cols = set()
    for c in columns:
        if c.field_name in record_keys or to_property_name(c.field_name, c.cpp_type) in record_keys or c.field_name == "RowName":
            matched_excel_cols.add(c.field_name)
    excel_columns_without_json = sorted(column_names - matched_excel_cols)
    if json_extra_fields:
        result.drift_issues.append(f"JSON 存在 Excel 无对应列的字段: {', '.join(json_extra_fields)}")
    if excel_columns_without_json:
        result.drift_issues.append(f"Excel 存在但 JSON 缺失的字段列: {', '.join(excel_columns_without_json)}")

    for record in json_map.values():
        for val in record.values():
            if isinstance(val, str) and val.strip().startswith(table_schema.TEXT_LITERAL_PREFIXES):
                result.text_literal_count += 1
    if result.text_literal_count > 0:
        result.warnings.append(
            f"[非派生 JSON] 本表 JSON 含 {result.text_literal_count} 处 FText 字面量（NSLOCTEXT/LOCTEXT），"
            f"应为 UE 侧导出物而非 Excel 正向派生；本次回写将按 Excel 可读形态归一，属预期规范化。")

    rule_records = [(seq, str(rec.get("RowName") or rec.get("Name") or "").strip(),
                     {c.field_name: rec.get(c.field_name) for c in columns})
                    for seq, rec in enumerate(json_records, start=1) if isinstance(rec, dict)]
    result.rule_issues.extend(table_rules.check_row_keys(rule_records, table_name))

    if options.strict and (result.conversion_errors or result.rule_issues or result.drift_issues):
        result.error = (f"前置数据校验失败（类型违规 {len(result.conversion_errors)} 处 / 规则违规 "
                        f"{len(result.rule_issues)} 处 / 结构漂移 {len(result.drift_issues)} 处，已终止写入）")
        return result

    rows_to_delete = []
    changes_count = 0

    for key, row_num in excel_rows.items():
        if key not in json_map:
            if options.allow_delete:
                rows_to_delete.append(row_num)
                result.diff_list.append({
                    "type": "DELETE_ROW",
                    "key": key,
                    "field": "*",
                    "old": f"Row {row_num}",
                    "new": "<DELETED>",
                    "row": row_num,
                    "col": 0,
                })
                changes_count += 1
            else:
                result.warnings.append(f"行 [{key}] 在 JSON 中缺失，已安全保留在 Excel 中（如需删除请指定 --allow-delete）。")
            continue

        record = json_map[key]
        for col in columns:
            prop_name = to_property_name(col.field_name, col.cpp_type)
            if col.field_name in record:
                raw_json_val = record[col.field_name]
            elif prop_name in record:
                raw_json_val = record[prop_name]
            elif col.field_name == "RowName" and "Name" in record:
                raw_json_val = record["Name"]
            else:
                continue

            cell = sheet.cell(row=row_num, column=col.col_idx + 1)
            orig_val = cell.value

            if orig_val is not None and str(orig_val).strip().startswith("="):
                if protect_formulas:
                    result.warnings.append(f"[{key}.{col.field_name}] 单元格为公式 ({orig_val})，已自动保护跳过。若需覆盖请用 --force-formula-overwrite")
                    continue

            if (key, col.field_name) in invalid_cells:
                continue

            try:
                target_val = table_schema.convert_for_excel(raw_json_val, col.cpp_type, fill_defaults=fill_defaults, strict=False)
            except table_schema.SchemaValidationError as e:
                result.conversion_errors.append(f"[{key}.{col.field_name}] {e}")
                continue

            if target_val is None and (orig_val is None or orig_val == ""):
                continue

            if not table_schema.values_are_equal(orig_val, target_val, col.cpp_type, epsilon=float_eps):
                if target_val is None and not fill_defaults:
                    result.warnings.append(f"[{key}.{col.field_name}] JSON 中为空值，保留 Excel 现有值 '{orig_val}'（如需清零请加 --fill-defaults）")
                    continue

                result.diff_list.append({
                    "type": "MODIFY",
                    "key": key,
                    "field": col.field_name,
                    "old": orig_val,
                    "new": target_val,
                    "row": row_num,
                    "col": col.col_idx + 1,
                })
                changes_count += 1
                if not options.dry_run:
                    cell.value = target_val

    last_valid_row = max(excel_rows.values()) if excel_rows else 3
    sample_style_row = last_valid_row if last_valid_row >= 4 else None

    for key in json_keys_ordered:
        if key in excel_rows:
            continue

        new_row_idx = last_valid_row + 1
        record = json_map[key]

        result.diff_list.append({
            "type": "ADD_ROW",
            "key": key,
            "field": "*",
            "old": "<NONE>",
            "new": f"New Row {new_row_idx}",
            "row": new_row_idx,
            "col": 0,
        })
        changes_count += 1

        if options.dry_run:
            last_valid_row = new_row_idx
            continue

        if last_valid_row < sheet.max_row:
            sheet.insert_rows(new_row_idx, 1)

        last_valid_row = new_row_idx
        excel_rows[key] = new_row_idx

        if sample_style_row:
            for col in columns:
                src_cell = sheet.cell(row=sample_style_row, column=col.col_idx + 1)
                dst_cell = sheet.cell(row=new_row_idx, column=col.col_idx + 1)
                copy_cell_style(src_cell, dst_cell)

        for col in columns:
            prop_name = to_property_name(col.field_name, col.cpp_type)
            raw_val = record.get(col.field_name)
            if raw_val is None and prop_name in record:
                raw_val = record.get(prop_name)
            if raw_val is None and col.field_name == "RowName":
                raw_val = key
            if (key, col.field_name) in invalid_cells:
                continue
            try:
                sheet.cell(row=new_row_idx, column=col.col_idx + 1).value = table_schema.convert_for_excel(
                    raw_val, col.cpp_type, fill_defaults=fill_defaults, strict=False
                )
            except table_schema.SchemaValidationError as e:
                result.conversion_errors.append(f"[{key}.{col.field_name}] 新增行 {e}")

    if rows_to_delete and not options.dry_run:
        if has_formula_in_data:
            formula_message = "[公式安全告警] 本表包含计算公式；执行物理删行可能导致外部公式引用位移或变为 #REF!。"
            if options.strict:
                result.error = f"{formula_message}（--strict 下已阻断删行，未做任何修改）"
                return result
            result.warnings.append(formula_message)
        for r in sorted(rows_to_delete, reverse=True):
            sheet.delete_rows(r, 1)

    if changes_count > 0 and not options.dry_run:
        if backup_session is not None:
            try:
                result.backup_path = backup_session.backup(xlsx_path)
            except OSError as e:
                result.error = f"写前备份原文件失败（中止写入）{xlsx_path}: {e}"
                return result

        tmp_path = xlsx_path + TMP_SUFFIX
        try:
            wb.save(tmp_path)
            os.replace(tmp_path, xlsx_path)
        except OSError as e:
            if os.path.isfile(tmp_path):
                try:
                    os.remove(tmp_path)
                except OSError as cleanup_error:
                    print(f"[警告] 无法清理临时文件 {tmp_path}: {cleanup_error}", file=sys.stderr)
            result.error = f"保存 Excel 失败（原文件未被破坏）{xlsx_path}: {e}"
            return result

    result.has_changes = changes_count > 0
    return result


def sync_all(cfg_path: Optional[str], options: Any) -> bool:
    """全量或定向执行 JSON ➔ Excel 反向同步"""
    cfg = tool_config.load_config(cfg_path, (tool_config.KEY_DATA_TABLES_SOURCE_DIR, "json_output_dir"))
    excel_dir = cfg[tool_config.KEY_DATA_TABLES_SOURCE_DIR]
    json_dir = cfg["json_output_dir"]
    backup_dir = resolve_backup_dir(cfg, cfg_path)
    max_backups = int(cfg.get("max_backups", 10))

    options.protect_formulas_config = cfg.get("protect_formulas", True)

    print("=" * 65)
    print(" JSON ➔ Excel 反向同步管线启动（单一真源 · 写前快照 · 原子落盘）")
    print(f" Excel 源目录 : {excel_dir}")
    print(f" JSON 数据目录: {json_dir}")
    print(f" 运行模式     : {'[DRY-RUN 仅预览]' if options.dry_run else '[正式回写]'}")
    print(f" 备份目录     : {backup_dir if not options.no_backup else '[已禁用备份]'}")
    print(f" 公式保护     : {'[强制覆盖]' if options.force_formula else ('[启用]' if options.protect_formulas_config else '[禁用]')}")
    print(f" 空值策略     : {'[回填类型默认值]' if options.fill_defaults else '[保持原单元格留空]'}")
    print(f" 浮点容差     : {options.float_epsilon}")
    print("=" * 65)

    matched_pairs = []
    for root, _, files in os.walk(json_dir):
        for file in files:
            if not file.endswith(".json") or file.startswith("."):
                continue

            json_path = os.path.join(root, file)
            rel_dir = os.path.relpath(root, json_dir)
            table_name = os.path.splitext(file)[0]

            if options.table:
                target_table = options.table.replace("DT_", "").strip()
                cur_table = table_name.replace("DT_", "").strip()
                if options.table != table_name and target_table != cur_table:
                    continue

            xlsx_path = os.path.normpath(os.path.join(excel_dir, rel_dir, f"{table_name}.xlsx"))
            if not os.path.isfile(xlsx_path):
                root_xlsx_path = os.path.normpath(os.path.join(excel_dir, f"{table_name}.xlsx"))
                if os.path.isfile(root_xlsx_path):
                    xlsx_path = root_xlsx_path
                else:
                    continue

            matched_pairs.append((xlsx_path, json_path, table_name))

    matched_pairs.sort(key=lambda item: item[2])
    if not matched_pairs:
        print(f"[提示] 未找到匹配的 JSON ➔ Excel 数据表对（搜索目录: {json_dir}）。")
        return True

    backup_session = BackupSession(backup_dir, excel_dir, max_backups) if not options.no_backup and not options.dry_run else None

    total_changed_tables = 0
    total_diff_count = 0
    total_conversion_errors = 0
    total_drift_issues = 0
    total_rule_issues = 0
    has_error = False

    for xlsx_path, json_path, table_name in matched_pairs:
        result = perform_sync_for_table(xlsx_path, json_path, options, backup_session)

        for w in result.warnings:
            print(f"  \033[33m[告警] {w}\033[0m")
        for ce in result.conversion_errors:
            print(f"  \033[33m[转换失败] {ce}\033[0m")
        for drift in result.drift_issues:
            print(f"  \033[33m[结构漂移] {drift}\033[0m")
        for issue in result.rule_issues:
            print(f"  \033[31m[规则违规] {issue.code} 主键[{issue.key or '-'}] 第 {issue.row_idx} 行\033[0m")
        total_conversion_errors += len(result.conversion_errors)
        total_drift_issues += len(result.drift_issues)
        total_rule_issues += len(result.rule_issues)

        if result.error:
            print(f"  \033[31m[错误] {result.error}\033[0m")
            has_error = True
            continue

        if result.has_changes:
            total_changed_tables += 1
            total_diff_count += len(result.diff_list)
            print(f"\033[36m▶ [{table_name}] 发现 {len(result.diff_list)} 处变动:\033[0m")
            for diff in result.diff_list[:15]:
                if diff["type"] == "MODIFY":
                    print(f"    • {diff['key']}.{diff['field']}: {diff['old']} ➔ \033[32m{diff['new']}\033[0m")
                elif diff["type"] == "ADD_ROW":
                    print(f"    • \033[32m[+新增行]\033[0m {diff['key']}")
                elif diff["type"] == "DELETE_ROW":
                    print(f"    • \033[31m[-删除行]\033[0m {diff['key']}")
            if len(result.diff_list) > 15:
                print(f"    ... 以及其余 {len(result.diff_list) - 15} 处变更")

            status_text = "[已预览]" if options.dry_run else "[✓ 已更新]"
            backup_note = f"（写前快照: {os.path.basename(result.backup_path)}）" if result.backup_path else ""
            print(f"  {status_text} {os.path.basename(xlsx_path)}{backup_note}\n")
        else:
            print(f"  [无变动] {table_name}")

    snapshot_dir = backup_session.finalize() if backup_session is not None else None

    print("=" * 65)
    print(f" 同步完成！共检测 {len(matched_pairs)} 张表，{total_changed_tables} 张表产生变动（共 {total_diff_count} 处变更）。")
    if total_conversion_errors > 0:
        print(f" \033[33m[转换失败汇总] 共 {total_conversion_errors} 处（已跳过、未写入；--strict 时计入退出码）\033[0m")
    if total_rule_issues > 0:
        print(f" \033[33m[规则违规汇总] 共 {total_rule_issues} 处（主键缺失/非法/重复；--strict 时计入退出码）\033[0m")
    if total_drift_issues > 0:
        print(f" \033[33m[结构漂移汇总] 共 {total_drift_issues} 处（JSON 与 Excel 表头不同步；--strict 时计入退出码）\033[0m")
    if snapshot_dir:
        print(f" \033[32m[写前快照已归档] 快照位置: {snapshot_dir}\033[0m")
        print(f" \033[32m[回滚方式] python3 json_to_excel.py --config <配置> --rollback\033[0m")
    elif options.dry_run:
        print(" \033[33m[Dry Run] 仅完成差异分析，未向磁盘写入任何更改。\033[0m")
    elif options.no_backup:
        print(" \033[33m[提示] 本次以 --no-backup 运行，未生成写前快照。\033[0m")
    print("=" * 65)

    if has_error:
        return False
    if options.strict and (total_conversion_errors or total_rule_issues or total_drift_issues) > 0:
        return False
    return True


def main() -> None:
    parser = argparse.ArgumentParser(description="JSON 反写 Excel 数据表同步工具（单一真源 · 写前快照 · 原子落盘）")
    parser.add_argument("--config", dest="config_path", default=None, help="JSON 配置文件路径")
    parser.add_argument("--table", dest="table", default=None, help="仅同步指定数据表名（如 DT_Example_Table）")
    parser.add_argument("--dry-run", dest="dry_run", action="store_true", help="仅打印变动差异报告，不修改 Excel 文件")
    parser.add_argument("--allow-delete", dest="allow_delete", action="store_true", help="允许物理删除 Excel 中存在但 JSON 中缺失的行")
    parser.add_argument("--force-formula-overwrite", dest="force_formula", action="store_true", help="强制覆盖带有计算公式的单元格")
    parser.add_argument("--fill-defaults", dest="fill_defaults", action="store_true", help="空值/null 写入对应 C++ 类型的默认零值（0/False）而不是保持 Excel 留空")
    parser.add_argument("--float-epsilon", dest="float_epsilon", type=float, default=table_schema.DEFAULT_FLOAT_EPSILON, help=f"浮点比较容差（默认 {table_schema.DEFAULT_FLOAT_EPSILON}）")
    parser.add_argument("--no-backup", dest="no_backup", action="store_true", help="禁用写前自动快照备份")
    parser.add_argument("--strict", dest="strict", action="store_true", help="存在转换失败时以非零退出码结束（供 CI 门使用）")
    parser.add_argument("--rollback", dest="rollback", action="store_true", help="按快照清单把写前文件还原回 Excel（不写回数据）")
    parser.add_argument("--snapshot", dest="snapshot", default=None, help="回滚时指定快照目录（缺省取最新时间戳快照）")
    parsed_args = parser.parse_args()

    if parsed_args.rollback:
        cfg = tool_config.load_config(parsed_args.config_path, (tool_config.KEY_DATA_TABLES_SOURCE_DIR,))
        backup_dir = resolve_backup_dir(cfg, parsed_args.config_path)
        ok, message = perform_rollback(backup_dir, parsed_args.snapshot, parsed_args.table)
        print(f"[回滚] {message}")
        sys.exit(0 if ok else 1)

    success = sync_all(parsed_args.config_path, parsed_args)
    sys.exit(0 if success else 1)


if __name__ == "__main__":
    main()
